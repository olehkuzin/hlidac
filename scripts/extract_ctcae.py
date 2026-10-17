"""Transcribe the demo CTCAE rows from the official NCI xlsx into data/ctcae/ (WP2).

Mechanical on purpose: every clinical string in the output is a cell value from the
NCI file with only leading/trailing whitespace stripped. Nothing is typed, merged or
reworded here. Metadata (title, publication date) is passed on the command line and
comes from the cover of NCI's own quick-reference PDF for the same version.

Run (openpyxl is a one-off dependency, not part of the project):

    uv run --with openpyxl python -I scripts/extract_ctcae.py CTCAE_v5.0.xlsx 5.0 \
        data/ctcae/ctcae-v5.0.yaml --source-url <url> --title <title> --published <date>

Treat the download as untrusted: keep it outside the repo and run with -I.
"""

import argparse
import hashlib
import pathlib
import re
import sys

import openpyxl
import yaml

TERMS = ("Fever", "Chills", "Nausea", "Vomiting", "Diarrhea")
APPROVAL = "DEMO — čeká na schválení FN Motol"
PUBLISHER = "National Cancer Institute (NCI), CTEP"
RETRIEVED = "2026-10-16"

# Column header patterns (the headers differ slightly between v5.0 and v6.0).
COLUMNS = {
    "code": r"(MedDRA .*Code|LLT Code)$",
    "soc": r"SOC$",
    "term": r"Term",
    "1": r"Grade 1$",
    "2": r"Grade 2$",
    "3": r"Grade 3$",
    "4": r"Grade 4$",
    "5": r"Grade 5$",
    "definition": r"Definition$",
    "navigational_note": r"Navigational Note$",
    "change": r"Change$",
}


def cell(value: object) -> str | None:
    """A cell as text, stripped; None for an empty cell. Never rewords anything."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def find_columns(header: tuple) -> dict[str, int]:
    found: dict[str, int] = {}
    for key, pattern in COLUMNS.items():
        hits = [i for i, h in enumerate(header) if h and re.search(pattern, str(h).strip())]
        if len(hits) != 1:
            sys.exit(f"column {key!r}: expected one header match, got {hits} in {header}")
        found[key] = hits[0]
    return found


def clean_sheet(wb: openpyxl.Workbook, version: str) -> str:
    names = [n for n in wb.sheetnames if f"v{version}" in n and "Clean Copy" in n]
    if len(names) != 1:
        sys.exit(f"expected one 'CTCAE v{version} Clean Copy' sheet, got {wb.sheetnames}")
    return names[0]


def errata_for(wb: openpyxl.Workbook, term: str) -> list[dict]:
    """Rows from NCI's own Errata sheet (v6.0 has one), verbatim."""
    if "Errata" not in wb.sheetnames:
        return []
    rows = list(wb["Errata"].iter_rows(values_only=True))
    header = [cell(h) for h in rows[0]]
    out = []
    for row in rows[1:]:
        rec = dict(zip(header, row, strict=False))
        if cell(rec.get(header[2])) != term:
            continue
        date = rec.get("Update Date")
        out.append({
            "item": cell(rec.get("Grade/Definition")),
            "previous": cell(rec.get("Previous Version")),
            "updated": cell(rec.get("Updated Version")),
            "update_date": date.date().isoformat() if hasattr(date, "date") else cell(date),
        })
    return out


def extract(xlsx: pathlib.Path, version: str) -> tuple[str, list[dict]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    sheet = clean_sheet(wb, version)
    rows = list(wb[sheet].iter_rows(values_only=True))
    col = find_columns(rows[0])
    by_term: dict[str, list[tuple]] = {}
    for row in rows[1:]:
        name = cell(row[col["term"]])
        if name in TERMS:
            by_term.setdefault(name, []).append(row)
    terms = []
    for name in TERMS:
        hits = by_term.get(name, [])
        if len(hits) != 1:
            sys.exit(f"{name!r}: expected exactly one row in {sheet!r}, got {len(hits)}")
        row = hits[0]
        code = row[col["code"]]
        terms.append({
            "term": cell(row[col["term"]]),
            "meddra_code": str(code) if isinstance(code, int) else cell(code),
            "soc": cell(row[col["soc"]]),
            "grades": {g: cell(row[col[g]]) for g in ("1", "2", "3", "4", "5")},
            "definition": cell(row[col["definition"]]),
            "navigational_note": cell(row[col["navigational_note"]]),
            "ctcae_change": cell(row[col["change"]]),
            "errata": errata_for(wb, name),
        })
    wb.close()
    return sheet, terms


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("xlsx", type=pathlib.Path)
    p.add_argument("version", help='e.g. "5.0"')
    p.add_argument("out", type=pathlib.Path)
    p.add_argument("--source-url", required=True)
    p.add_argument("--title", required=True, help="from the cover of NCI's quick-reference PDF")
    p.add_argument("--published", default=None, help="publication date as NCI states it")
    a = p.parse_args()

    sheet, terms = extract(a.xlsx, a.version)
    doc = {
        "content_id": f"ctcae-v{a.version}",
        "ctcae_version": a.version,
        "approval": APPROVAL,
        "source": {
            "publisher": PUBLISHER,
            "title": a.title,
            "url": a.source_url,
            "file": a.xlsx.name,
            "sheet": sheet,
            "sha256": hashlib.sha256(a.xlsx.read_bytes()).hexdigest(),
            "retrieved": RETRIEVED,
            "published": a.published,
        },
        "terms": terms,
    }
    head = (
        "# GENERATED by scripts/extract_ctcae.py from the official NCI CTCAE xlsx. Do not edit.\n"
        "# Every string under `terms` is a verbatim cell value (only outer whitespace stripped).\n"
        "# `errata` rows come verbatim from NCI's Errata sheet in the same file.\n"
    )
    body = yaml.safe_dump(doc, allow_unicode=True, sort_keys=False, width=10_000)
    a.out.parent.mkdir(parents=True, exist_ok=True)
    a.out.write_text(head + body, encoding="utf-8")
    print(f"wrote {a.out} ({len(terms)} terms from {sheet!r})")


if __name__ == "__main__":
    main()
